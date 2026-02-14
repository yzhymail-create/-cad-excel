from __future__ import annotations

import logging
from pathlib import Path


logger = logging.getLogger(__name__)


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    if suffix == ".xlsx":
        return _extract_xlsx(path)
    if suffix == ".ppt":
        logger.warning(
            "Legacy .ppt format detected (%s). File will be skipped; convert to .pptx using PowerPoint or LibreOffice.",
            path,
        )
        return ""
    if suffix == ".pptx":
        return _extract_pptx(path)
    if suffix in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
        return _extract_image(path)
    logger.warning("Unsupported file extension: %s", path)
    return ""


def _extract_pdf(path: Path) -> str:
    from pdfminer.high_level import extract_text as pdf_extract_text

    return pdf_extract_text(str(path))


def _extract_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    return "\n".join(p.text for p in document.paragraphs if p.text)


def _extract_xlsx(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(filename=str(path), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            values = [
                str(cell).strip()
                for cell in row
                if cell is not None and str(cell).strip()
            ]
            if values:
                parts.append(" ".join(values))
    workbook.close()
    return "\n".join(parts)


def _extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            text = getattr(shape, "text", "")
            if text:
                parts.append(text)
    return "\n".join(parts)


def _extract_image(path: Path) -> str:
    try:
        from PIL import Image
        import pytesseract
    except ImportError:
        logger.warning("OCR dependencies missing; skipping image %s", path)
        return ""

    try:
        return pytesseract.image_to_string(Image.open(path))
    except Exception as exc:  # noqa: BLE001 - third-party errors vary
        logger.warning("OCR failed for %s: %s", path, exc)
        return ""
